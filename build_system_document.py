from pathlib import Path
import sqlite3
from datetime import date

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "安嘉之越南美食品在淡水_系統文件書_最新版.docx"
DB = ROOT / "vietnam_food.db"
FIGURES = [
    (ROOT / "system-diagrams" / "01_整體系統架構圖.png", "圖 3-1　整體系統架構圖"),
    (ROOT / "system-diagrams" / "02_使用案例圖.png", "圖 3-2　使用案例圖"),
    (ROOT / "system-diagrams" / "03_登入畫面流程圖.png", "圖 3-3　登入畫面流程圖"),
    (ROOT / "system-diagrams" / "04_領域類別圖.png", "圖 3-4　領域類別圖"),
]

NAVY = "17365D"
TEAL = "197A68"
GOLD = "B77B1C"
LIGHT = "EAF2F6"
GRAY = "5B6573"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, val in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def set_table_widths(table, widths):
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = Inches(width)
            tc_pr = row.cells[idx]._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width * 1440)))
            tc_w.set(qn("w:type"), "dxa")


def set_font(run, size=12, bold=False, color="000000", italic=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "標楷體")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def add_run(p, text, size=12, bold=False, color="000000", italic=False):
    r = p.add_run(text)
    set_font(r, size=size, bold=bold, color=color, italic=italic)
    return r


def format_paragraph(p, before=0, after=6, line=1.0, align=None):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    if align is not None:
        p.alignment = align


def add_para(doc, text="", size=12, bold=False, color="000000", italic=False, before=0, after=6, align=None):
    p = doc.add_paragraph()
    format_paragraph(p, before, after, 1.0, align)
    add_run(p, text, size, bold, color, italic)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    sizes = {1: 16, 2: 14, 3: 12}
    color = NAVY if level == 1 else TEAL
    format_paragraph(p, before=14 if level == 1 else 10, after=7 if level == 1 else 5)
    add_run(p, text, sizes[level], True, color)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    format_paragraph(p, after=3, line=1.0)
    add_run(p, text, 12)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    format_paragraph(p, after=3, line=1.0)
    add_run(p, text, 12)
    return p


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        format_paragraph(p, after=0, align=WD_ALIGN_PARAGRAPH.CENTER)
        add_run(p, h, 11, True, "FFFFFF")
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_margins(cells[i])
            if ridx % 2 == 1:
                set_cell_shading(cells[i], "F7FAFC")
            p = cells[i].paragraphs[0]
            format_paragraph(p, after=0)
            # The department format requests 12 pt body text; tables retain the
            # same body size and rely on row wrapping instead of shrinking text.
            add_run(p, str(value), 12)
    if widths:
        set_table_widths(table, widths)
    add_para(doc, "", size=4, after=1)
    return table


def page_number(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    set_font(run, 10, color=GRAY)


def setup_document():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.9)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "標楷體")
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.0
    normal.paragraph_format.space_after = Pt(6)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    format_paragraph(header, after=0)
    add_run(header, "安嘉之越南美食品在淡水｜系統文件書", 9.5, False, GRAY)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    format_paragraph(footer, after=0)
    add_run(footer, "淡江大學資訊管理學系　", 10, False, GRAY)
    page_number(footer)
    return doc


def database_summary():
    fallback = {"menu": 0, "orders": 0, "order_items": 0, "accounts": 0}
    if not DB.exists():
        return fallback
    try:
        with sqlite3.connect(DB) as conn:
            return {
                "menu": conn.execute("SELECT COUNT(*) FROM menu").fetchone()[0],
                "orders": conn.execute("SELECT COUNT(*) FROM orders").fetchone()[0],
                "order_items": conn.execute("SELECT COUNT(*) FROM order_items").fetchone()[0],
                "accounts": conn.execute("SELECT COUNT(*) FROM customer_accounts").fetchone()[0],
            }
    except Exception:
        return fallback


def add_cover(doc):
    for _ in range(6):
        add_para(doc, "", after=0)
    add_para(doc, "淡江大學資訊管理學系", 18, True, NAVY, after=14, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_para(doc, "畢業專題系統文件書", 22, True, NAVY, after=34, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_para(doc, "安嘉之越南美食品在淡水", 26, True, TEAL, after=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_para(doc, "Ẩm thực Việt Nam tại Tamsui – Thương hiệu Tuấn An và Dịch Gia", 13, True, GOLD, after=30, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_para(doc, "三店整合式行動外帶點餐與營運管理系統", 15, True, NAVY, after=42, align=WD_ALIGN_PARAGRAPH.CENTER)
    meta = [
        ("專題組員", "楊竣安　411631053　資管 4C\n陳奕嘉　411631384　資管 4C"),
        ("指導老師", "周清江老師"),
        ("文件版本", "最新版（依 2026 年 8 月系統功能整理）"),
        ("完成日期", date.today().strftime("%Y 年 %m 月 %d 日")),
    ]
    add_table(doc, ["項目", "內容"], meta, [1.5, 4.7])
    add_para(doc, "本文件以目前可運作之 Flask、SQLite、Cloudflare Tunnel 外帶點餐系統為範圍，說明設計、實作與驗證成果。", 11, False, GRAY, before=30, after=0, align=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_page_break()


def add_toc(doc):
    add_heading(doc, "目錄", 1)
    entries = [
        "摘要", "第一章　緒論", "第二章　相關技術應用與重要文獻", "第三章　系統概要設計", "第四章　系統開發工具與使用環境",
        "第五章　系統實作及實驗結果", "第六章　結論及未來發展", "參考文獻", "附錄一　工作分配與專案進度甘特圖表",
        "附錄二　資料庫資料表與欄位說明"
    ]
    for index, entry in enumerate(entries, 1):
        p = doc.add_paragraph()
        format_paragraph(p, after=4)
        add_run(p, entry, 12)
        add_run(p, "　" + str(index + 1), 12, False, GRAY)
    add_para(doc, "說明：頁碼將依 Word 開啟後的自動版面配置而變動；本目錄提供閱讀導覽。", 10.5, False, GRAY, before=10)
    doc.add_page_break()


def add_summary(doc, stats):
    add_heading(doc, "摘要", 1)
    add_para(doc, "本專題以淡水地區越南美食外帶需求為背景，建置「安嘉之越南美食品在淡水」行動點餐與營運管理系統。系統整合安嘉河食、安嘉越南麵包與安嘉越南食世界三間店家，使用者可透過手機瀏覽中越雙語菜單、選擇客製加料、以現金付款、註冊會員、累積點數、查詢當日外帶訂單與取餐流水號；店家管理者可即時處理訂單、調整供應狀態、維護菜單與圖片、人工點餐、查看每日銷量及歷史營運報表。")
    add_para(doc, "系統採用 Python Flask 建立 Web 應用程式，以 SQLite 儲存菜單、訂單、訂單明細、會員帳號與系統設定，並透過 Cloudflare Tunnel 將本機服務以 HTTPS 公開於網際網路。介面採行動裝置優先設計，兼顧中文與越南文的資訊呈現，並以分類、推薦標示、地區／起源與供應狀態協助使用者快速完成外帶點餐。")
    add_para(doc, f"截至本文件產製時，資料庫已保存 {stats['menu']} 筆菜單資料、{stats['orders']} 筆訂單、{stats['order_items']} 筆訂單明細與 {stats['accounts']} 筆會員帳號資料；實際數量會隨營運持續更新。系統成果顯示，小型餐飲店可使用低成本的 Web 技術完成多店整合、即時訂單處理與基本營運分析。")
    add_para(doc, "關鍵詞：越南美食、外帶點餐、Flask、SQLite、會員點數、Cloudflare Tunnel、中越雙語", 12, True, TEAL, before=8)
    doc.add_page_break()


def add_chapter1(doc):
    add_heading(doc, "第一章　緒論", 1)
    add_heading(doc, "1.1 研究背景", 2)
    add_para(doc, "外帶餐飲服務需要在短時間內完成菜單瀏覽、客製選項、訂單確認、叫號與取餐。傳統紙本菜單與口頭點餐容易造成餐點名稱、加料、付款或製作狀態的溝通落差；而越南料理常含有中越雙語名稱、南北用語差異與多種配料，資訊呈現需求更為明顯。本專題因此以淡水地區的越南美食品牌為情境，開發適合手機瀏覽的外帶點餐系統。")
    add_heading(doc, "1.2 專題目的與重要性", 2)
    for t in [
        "建立三間店家共用的外帶點餐入口，減少顧客切換不同系統的成本。",
        "提供繁體中文與越南文菜單、食材介紹與用語輔助，降低跨語言點餐門檻。",
        "以即時訂單刷新、取餐流水號及製作狀態，協助現場店員快速掌握訂單。",
        "提供後台菜單、圖片、加料、供應狀態、會員點數與營運報表管理功能。",
        "保留可擴充架構，未來可串接電子支付、雲端資料庫與正式部署環境。",
    ]:
        add_bullet(doc, t)
    add_heading(doc, "1.3 主要研究問題與目標", 2)
    add_table(doc, ["研究問題", "系統回應方式"], [
        ("多店菜單如何整合", "以 shop 欄位與三個店別頁面分流，後台亦可依店別管理。"),
        ("客製加料如何正確記錄", "以每筆訂單明細保存餐點、加料、備註、單價與數量，避免僅存單一訂單文字。"),
        ("顧客如何得知取餐進度", "建立已接單、製作中、已完成可領取、已領取四階段製作狀態與每日店別流水號。"),
        ("如何鼓勵會員回購", "提供註冊、歡迎點數、消費累積點數、餐點兌換與金額折抵設定。"),
        ("如何以低成本公開服務", "使用 Cloudflare Tunnel 建立 HTTPS 公開網址，讓手機可透過行動網路連線。"),
    ], [2.1, 4.1])
    add_heading(doc, "1.4 系統使用對象", 2)
    add_table(doc, ["角色", "使用情境", "主要權限／功能"], [
        ("一般顧客", "以手機瀏覽菜單並完成外帶點餐", "選店、選餐、加料、備註、現金付款、查詢訂單與流水號"),
        ("註冊會員", "累積並使用點數", "帳號登入、領取歡迎點數、累積消費點數、餐點兌換、金額折抵"),
        ("店家管理者", "處理現場與線上外帶訂單", "訂單狀態與付款狀態、人工點餐、叫號、營運報表"),
        ("菜單維護人員", "更新菜色及供應資訊", "新增、編輯、下架商品，更新圖片、分類、加料、供應狀態與推薦標示"),
    ], [1.3, 2.3, 2.6])
    add_heading(doc, "1.5 系統特色", 2)
    add_para(doc, "本系統的特色可歸納為「三店整合、雙語導覽、外帶流程、即時管理、會員回饋、可視化報表」六項。不同於一般單店菜單，本系統將河粉、法國麵包、飲品點心整合至同一品牌入口；同時讓管理者可依店別處理訂單與查看銷量，避免混單。")
    doc.add_page_break()


def add_chapter2(doc):
    add_heading(doc, "第二章　相關技術應用與重要文獻", 1)
    add_heading(doc, "2.1 Web 應用架構", 2)
    add_para(doc, "系統採用 Flask 作為後端 Web 框架，負責路由處理、表單驗證、模板渲染與資料庫交易。Flask 結構輕量、學習成本低，適合畢業專題以迭代方式逐步新增會員、點數、報表及後台維護等功能。前端以 HTML、CSS 與 JavaScript 實作響應式介面與定時同步機制。")
    add_heading(doc, "2.2 SQLite 與交易資料保存", 2)
    add_para(doc, "SQLite 為嵌入式關聯式資料庫，將資料保存於單一資料庫檔案，適合小型餐飲資訊系統的開發、展示與單機營運。系統將菜單資料、客製加料、訂單主檔、訂單明細、會員帳號、密碼重設權杖及系統設定分表保存，確保新增餐點、下架商品與歷史訂單可以被追溯。")
    add_heading(doc, "2.3 HTTPS 公開存取與雲端通道", 2)
    add_para(doc, "系統透過 Cloudflare Tunnel 將本機 Flask 服務安全地連接至公開網域，讓消費者可使用 HTTPS 網址掃描 QR Code 後開啟點餐頁面。此作法降低自行設定公網 IP、路由器轉發與憑證維護的複雜度，但仍需保持承載主機與 Tunnel 程序運作。")
    add_heading(doc, "2.4 行動裝置優先與可用性設計", 2)
    add_para(doc, "介面以手機螢幕為主要情境，使用大型按鈕、清楚的區塊階層與繁中／越南文並列資訊。菜單卡片顯示圖片、名稱、介紹、地區／起源、推薦標示及供應狀態；購物車則顯示原價、折抵、實付金額與取餐資訊，協助顧客在網路環境與時間壓力下快速確認。")
    add_heading(doc, "2.5 與既有作法比較", 2)
    add_table(doc, ["面向", "紙本／口頭點餐", "一般線上菜單", "本系統"], [
        ("多店管理", "需分別製作與整理", "通常單店為主", "三店共用入口並可店別管理"),
        ("雙語資訊", "容易受版面限制", "多為單一語言", "繁中、越南文與南北用語提示"),
        ("客製加料", "口頭紀錄易遺漏", "選項固定", "可針對每道餐點設定、顯示／隱藏與調整價格"),
        ("製作流程", "人工喊號與紙條", "未必有後台", "即時訂單、狀態、付款與每日店別流水號"),
        ("營運分析", "需另行彙整", "常限於平台資料", "每日銷量、歷史報表、折線圖與圓餅圖"),
    ], [1.25, 1.7, 1.7, 1.85])
    doc.add_page_break()


def add_chapter3(doc):
    add_heading(doc, "第三章　系統概要設計", 1)
    add_heading(doc, "3.1 開發方法與設計原則", 2)
    add_para(doc, "本專題採漸進式原型開發。先完成外帶點餐、訂單與後台，再依實際操作需求增補菜單圖片、雙語食材介紹、點數機制、人工點餐、供應狀態及報表。設計原則為：資料正確性優先、行動裝置優先、三店可分流、前後台資料同步，以及變更菜單後保留既有訂單歷史。")
    add_heading(doc, "3.2 系統處理流程", 2)
    for step in [
        "顧客進入使用者首頁，選擇三間店家之一，瀏覽分類菜單及供應狀態。",
        "顧客選擇餐點、數量、可用加料與備註，加入購物車；系統在前端顯示小計、點數折抵與實付金額。",
        "顧客輸入取餐姓名與電話末三碼識別資訊，確認一律現金付款後送出訂單。",
        "後端驗證餐點與加料是否仍供應，建立 orders 與 order_items 資料，計算每日店別流水號與會員點數。",
        "後台外帶訂單管理自動輪詢最新訂單；店員更新製作狀態與付款狀態，完成後依流水號叫號。",
        "營運人員在日銷量與歷史報表頁面依店別、日期查看訂單數、營收、餐點銷量與圖表。",
    ]:
        add_number(doc, step)
    add_heading(doc, "3.3 系統架構與使用案例", 2)
    for path, caption in FIGURES[:2]:
        if path.exists():
            doc.add_picture(str(path), width=Inches(6.1))
            p = doc.add_paragraph()
            format_paragraph(p, after=8, align=WD_ALIGN_PARAGRAPH.CENTER)
            add_run(p, caption, 10.5, True, GRAY)
    add_heading(doc, "3.4 登入流程與領域類別", 2)
    for path, caption in FIGURES[2:]:
        if path.exists():
            doc.add_picture(str(path), width=Inches(6.1))
            p = doc.add_paragraph()
            format_paragraph(p, after=8, align=WD_ALIGN_PARAGRAPH.CENTER)
            add_run(p, caption, 10.5, True, GRAY)
    add_heading(doc, "3.5 資料表關連與檔案關連", 2)
    add_table(doc, ["層次", "主要元件", "責任"], [
        ("展示層", "templates/*.html、static/styles.css、圖片資產", "提供使用者首頁、點餐、訂單、後台與報表介面。"),
        ("應用層", "app.py、Flask routes、表單驗證", "執行商業規則、點數計算、訂單建立、圖片處理與狀態更新。"),
        ("資料層", "vietnam_food.db（SQLite）", "保存菜單、帳號、訂單、明細、加料設定與系統設定。"),
        ("公開存取層", "Cloudflare Tunnel、HTTPS 網域", "讓行動網路與 QR Code 可安全連入本機 Web 服務。"),
    ], [1.25, 2.35, 2.9])
    doc.add_page_break()


def add_chapter4(doc):
    add_heading(doc, "第四章　系統開發工具與使用環境", 1)
    add_table(doc, ["類別", "工具／技術", "用途與選擇理由"], [
        ("程式語言", "Python 3.14", "語法易讀，適合快速開發與 SQLite、Flask 整合。"),
        ("Web 框架", "Flask", "提供路由、模板、請求處理與輕量化的伺服器端架構。"),
        ("資料庫", "SQLite", "單檔保存、易於備份與展示，符合小型系統原型需求。"),
        ("前端", "HTML5、CSS3、JavaScript", "建構 RWD 介面、購物車互動、狀態輪詢與圖表呈現。"),
        ("開發工具", "Visual Studio Code、PowerShell", "程式編輯、除錯、終端機啟動與版本維護。"),
        ("網路服務", "Cloudflare Tunnel、HTTPS", "提供公開網域、TLS 憑證與行動網路存取。"),
        ("繪圖與文件", "系統圖、Word", "製作架構、使用案例、流程與領域類別圖及專題文件。"),
    ], [1.25, 1.65, 3.6])
    add_heading(doc, "4.1 使用環境需求", 2)
    add_table(doc, ["項目", "需求"], [
        ("管理端電腦", "Windows 10/11、Python 3.14、專案虛擬環境與 Flask 套件。"),
        ("使用者裝置", "Android 或 iOS 手機、具 HTTPS 瀏覽器與網際網路。"),
        ("網路", "店家主機須保持連線；Cloudflare Tunnel 程序須正常運行。"),
        ("資料備份", "定期複製 vietnam_food.db，並備份 static/menu-images/uploads 圖片資料夾。"),
        ("安全建議", "正式使用時應設定後台權限、備份排程、環境變數與更完整的部署服務。"),
    ], [1.5, 5.0])
    add_heading(doc, "4.2 系統限制與風險控管", 2)
    add_para(doc, "目前系統以個人電腦執行 Flask 與 SQLite；若電腦關機、網路中斷、Flask 程序或 Cloudflare Tunnel 停止，公開網址將無法回應。因此系統適合作為原型與小型店家現場工具，若需 24 小時營運，應遷移至雲端主機與託管資料庫，並增加監控、備援與權限管理。")
    doc.add_page_break()


def add_chapter5(doc, stats):
    add_heading(doc, "第五章　系統實作及實驗結果", 1)
    add_heading(doc, "5.1 前端功能實作", 2)
    add_table(doc, ["功能", "實作結果"], [
        ("三店首頁", "使用者可從同一首頁進入安嘉河食、安嘉越南麵包、安嘉越南食世界。"),
        ("雙語菜單", "每道餐點呈現中文、越文、食材／介紹、地區／起源、推薦標示、圖片與供應狀態。"),
        ("外帶購物車", "支援多道餐點、個別加料、備註、數量及現金付款確認。"),
        ("會員與點數", "支援註冊、登入、歡迎點數、每消費 50 元無條件進位累積 1 點、餐點兌換及金額折抵。"),
        ("訂單查詢", "依電話查詢當日訂單、三家店分組資訊、取餐號與歷程紀錄。"),
    ], [1.6, 4.9])
    add_heading(doc, "5.2 後台功能實作", 2)
    add_table(doc, ["功能", "實作結果"], [
        ("菜單即時管理", "可分店新增、編輯、分類、上傳圖片、設定推薦、供應狀態、點數及刪除商品與圖片。"),
        ("加料管理", "可逐道餐點新增／刪除／改名／調整價格與設定顯示／隱藏；前端使用相同設定。"),
        ("外帶訂單管理", "依三店檢視當日訂單，顯示顧客、電話末三碼、餐點明細、原價、實付、付款與製作狀態。"),
        ("人工點餐", "店員可替未帶手機顧客建立訂單；會員電話可累積點數。"),
        ("營運報表", "分店日銷量、歷史報表、折線圖與圓餅圖，支援日期區間查詢。"),
        ("系統設定", "可設定餐點免費兌換點數價值與本次金額折抵點數價值；目前預設 1 點折抵 NT$1。"),
    ], [1.6, 4.9])
    add_heading(doc, "5.3 驗證方式與結果", 2)
    add_para(doc, "本專題以功能導向測試驗證流程，包含：三店菜單載入、已賣完商品不可加入購物車、多餐點與多加料訂單建立、每日店別流水號、會員註冊登入、點數累積／兌換／折抵、後台狀態與付款更新、菜單圖片更新／刪除、人工點餐及報表查詢。測試採 Flask 測試用戶端及實際瀏覽器操作；路由回應成功與資料寫入正確為主要驗收標準。")
    add_table(doc, ["驗收項目", "預期結果", "結果"], [
        ("外帶下單", "多項餐點與各自加料皆能建立訂單", "通過"),
        ("取餐流水號", "依當日、店別順序由 1 重新編號", "通過"),
        ("後台即時訂單", "新訂單可於輪詢週期內出現，無須手動重新整理", "通過"),
        ("供應狀態", "暫停供應商品不可下單，前端顯示狀態", "通過"),
        ("菜單圖片", "上傳後顯示指定圖片；刪除商品同步移除圖片檔案", "通過"),
        ("點數規則", "消費累積與金額折抵依後台設定計算", "通過"),
    ], [1.65, 3.8, 1.05])
    add_heading(doc, "5.4 成果評估", 2)
    add_para(doc, "在正確性方面，系統將訂單主檔與訂單明細分離，能保存每一品項、加料與價格；在完整性方面，涵蓋顧客、會員、後台、菜單、訂單、報表及點數流程；在穩定性方面，針對無效加料、停售商品與過期操作進行伺服器端驗證；在美感與實用性方面，採行動優先、雙語與圖片化菜單；在創新與擴充性方面，提供三店整合、店別流水號、可配置加料及點數設定。")
    add_heading(doc, "5.5 SWOT、STP 與 4P 分析", 2)
    add_table(doc, ["SWOT", "分析"], [
        ("Strengths 優勢", "三店整合、雙語內容、客製加料、即時後台、會員點數與低成本 HTTPS 公開。"),
        ("Weaknesses 劣勢", "目前依賴本機主機與 Tunnel；SQLite 不適合大量併發與長期雲端營運。"),
        ("Opportunities 機會", "淡水地區多元文化餐飲、行動點餐、外帶需求與越南文化導覽需求持續成長。"),
        ("Threats 威脅", "外送平台競爭、網路或主機中斷、食材供應變動與資料安全需求。"),
    ], [1.5, 5.0])
    add_table(doc, ["行銷面向", "分析與系統對應"], [
        ("STP：市場區隔", "以淡水學生、居民、越南新住民及喜愛東南亞料理的行動族群為主。"),
        ("STP：目標與定位", "定位為可用繁中／越文安心操作的三店外帶點餐品牌入口。"),
        ("4P：產品 Product", "河粉米線、法國麵包、越南咖啡、飲品點心與客製加料。"),
        ("4P：價格 Price", "清楚顯示原價、活動價、加料價格、點數換餐與點數折抵。"),
        ("4P：通路 Place", "手機瀏覽器、QR Code、HTTPS 公開網址及店內人工點餐。"),
        ("4P：推廣 Promotion", "老闆推薦、招牌標示、歡迎點數、每日流水號與文化導覽內容。"),
    ], [1.55, 4.95])
    doc.add_page_break()


def add_chapter6(doc):
    add_heading(doc, "第六章　結論及未來發展", 1)
    add_heading(doc, "6.1 結論", 2)
    add_para(doc, "本專題完成一套以手機外帶點餐為核心的三店整合系統，將中越雙語菜單、客製加料、會員點數、店別流水號、訂單製作狀態、菜單即時管理、人工點餐與營運分析整合於同一 Web 應用程式。系統以 SQLite 保存歷史資料，並以 Flask 和 Cloudflare Tunnel 讓使用者透過 HTTPS 網址進入，達成可展示、可操作且能回應店家日常需求的成果。")
    add_heading(doc, "6.2 遭遇問題與挑戰", 2)
    for t in [
        "本機主機、Flask 與 Tunnel 任一程序停止時，公開網址會出現 502，顯示部署可用性仍需改善。",
        "菜單圖片、資料庫記錄與檔案路徑需同步維護，否則可能出現圖片與菜名不符或圖片遺失。",
        "不同餐點的加料、冰塊、主食與特殊選項規則不同，需要以可設定的資料結構避免硬編碼。",
        "中越翻譯存在南北用語差異，應由店家審核菜名與食材用語，而不完全依賴自動翻譯。",
    ]:
        add_bullet(doc, t)
    add_heading(doc, "6.3 未來發展", 2)
    for t in [
        "將服務移轉至雲端主機，使用 PostgreSQL 或 MySQL、備份機制與監控服務，提供 24 小時營運。",
        "導入角色權限、管理者登入、多店員帳號與操作稽核紀錄，提高後台安全性。",
        "串接電子支付、電子發票、庫存警示、桌面叫號顯示器與簡訊／推播通知。",
        "增加多語言管理流程與人工審核介面，建立更完整的越南料理詞彙資料庫。",
        "擴充消費者行為分析、推薦系統、活動券與更精準的營運預測模型。",
    ]:
        add_bullet(doc, t)
    add_heading(doc, "6.4 專題貢獻", 2)
    add_para(doc, "本專題的主要貢獻在於將餐飲外帶的實際流程轉化為可操作的資訊系統，並將文化與語言需求納入介面設計。除了實現點餐交易，本系統也展示資料庫設計、Web 開發、部署、行動使用者經驗與管理報表的整合能力，具備後續擴充為正式小型餐飲管理服務的基礎。")
    doc.add_page_break()


def add_references(doc):
    add_heading(doc, "參考文獻", 1)
    refs = [
        "[1] Pallets. (2026). Flask Documentation. https://flask.palletsprojects.com/",
        "[2] SQLite. (2026). SQLite Documentation. https://www.sqlite.org/docs.html",
        "[3] Cloudflare. (2026). Cloudflare Tunnel Documentation. https://developers.cloudflare.com/cloudflare-one/connections/connect-networks/",
        "[4] MDN Web Docs. (2026). Responsive design. https://developer.mozilla.org/",
        "[5] Nielsen Norman Group. (2026). Usability and mobile user experience articles. https://www.nngroup.com/",
        "[6] OWASP Foundation. (2026). OWASP Top 10 Web Application Security Risks. https://owasp.org/www-project-top-ten/",
        "[7] 淡江大學資訊管理學系。系統文件書格式與繳交說明（依課程公告格式整理）。",
    ]
    for r in refs:
        add_number(doc, r)
    add_heading(doc, "附錄一　工作分配與專案進度甘特圖表", 1)
    add_table(doc, ["工作項目", "楊竣安", "陳奕嘉", "期間"], [
        ("需求訪談、情境規劃與菜單資料整理", "主要負責", "共同協作", "2026/05/13–2026/06/06"),
        ("前端頁面、雙語內容與手機介面", "主要負責", "共同協作", "2026/06/07–2026/07/10"),
        ("後端、資料庫、訂單與後台管理", "共同協作", "主要負責", "2026/06/15–2026/08/15"),
        ("會員、點數、報表、測試與修正", "共同協作", "主要負責", "2026/08/01–2026/08/28"),
        ("系統文件、簡報、展示與驗收", "主要負責", "共同協作", "2026/08/20–2026/09/11"),
    ], [2.6, 1.1, 1.1, 1.7])
    add_table(doc, ["階段", "05/13–06/06", "06/07–07/10", "07/11–08/15", "08/16–08/28", "08/29–09/11"], [
        ("需求／企劃", "■■■■", "", "", "", ""),
        ("UI／菜單與雙語", "■■", "■■■■", "", "", ""),
        ("後端／資料庫", "", "■■■", "■■■■", "", ""),
        ("測試／優化", "", "", "■■", "■■■■", ""),
        ("文件／簡報／驗收", "", "", "", "■■", "■■■■"),
    ], [1.4, 1.1, 1.3, 1.25, 1.25, 1.2])
    doc.add_page_break()


def add_appendix_db(doc):
    add_heading(doc, "附錄二　資料庫資料表與欄位說明", 1)
    add_table(doc, ["資料表", "用途", "代表欄位"], [
        ("menu", "保存三間店菜單、分類、雙語介紹、圖片、供應與點數設定", "id、name、vi_name、price、shop、category、image_url、availability_status"),
        ("orders", "保存每筆外帶訂單主檔與狀態", "id、customer_name、phone、total_price、payment_status、status、created_at、shop_name"),
        ("order_items", "保存一筆訂單的各餐點、加料與數量", "order_id、menu_item_id、item_name、addon_name、unit_price、quantity、line_total"),
        ("customer_accounts", "會員帳號、電話、密碼雜湊與點數", "display_name、phone、password_hash、points、welcome_points_claimed"),
        ("app_settings", "保存可調整之系統設定", "setting_key、setting_value（如點數兌換／折抵價值）"),
        ("retired_menu_items", "記錄已下架菜單，避免初始化時再次建立", "shop、vi_name、retired_at"),
        ("password_reset_tokens", "保留未來忘記密碼流程的權杖資料", "token、account_id、expires_at"),
    ], [1.55, 2.3, 2.85])
    add_heading(doc, "附錄三　重要操作說明", 1)
    for t in [
        "啟動本機服務：使用專案虛擬環境執行 app.py，確認 Flask 監聽 port 5000。",
        "公開網址：確認 Cloudflare Tunnel 為 Healthy，且 Published application route 指向 http://localhost:5000。",
        "資料備份：在關閉寫入中的程序或定期排程下，備份 vietnam_food.db 與 static/menu-images/uploads。",
        "菜單更新：在後台按儲存餐點後，前端於短時間輪詢更新；若圖片更換，請選檔並儲存以寫入新路徑。",
        "訂單處理：店員依當日店別流水號叫號，完成後將製作狀態改為已完成可領取或已領取，並確認付款狀態。",
    ]:
        add_number(doc, t)


def main():
    stats = database_summary()
    doc = setup_document()
    add_cover(doc)
    add_toc(doc)
    add_summary(doc, stats)
    add_chapter1(doc)
    add_chapter2(doc)
    add_chapter3(doc)
    add_chapter4(doc)
    add_chapter5(doc, stats)
    add_chapter6(doc)
    add_references(doc)
    add_appendix_db(doc)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
